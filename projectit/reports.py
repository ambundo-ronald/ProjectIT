"""Manager reports: how the maintenance work is going, and whether service contracts pay.

Everything is worked out from the jobs themselves for a chosen period (7, 30, 90 or 365 days):
jobs opened and completed, how long repairs take, preventive maintenance done on time, who did the
work, which machines keep breaking, requests received, and for each service contract the money in
(subscription invoices plus extra work billed) against the money out (parts at cost plus labour hours
at the internal cost per hour from Techy Ops Settings).

Labour hours are the time logged with the job timer, per person; jobs without timer entries fall back to
their recorded labour hours, shared between the assignees.

Jobs completed before start and finish times were recorded fall back to their last-modified time.

The technician scorecard weighs how well the work was done, not only how much: jobs finished by their due
date, repairs that held (the machine needed no other repair within 30 days), jobs that never went back
from review or completion, and repairs coded with what was wrong and why. Each person gets a score out of
100 from the measures that apply to their jobs, once they have completed enough jobs for it to mean
something.
"""
from collections import Counter, defaultdict
from datetime import timedelta

import frappe
from frappe import _
from frappe.utils import add_days, add_months, cint, flt, get_datetime, getdate, now_datetime, nowdate

from projectit import billing, contracts, ops, timer

PERIODS = (7, 30, 90, 365)
JOB_FIELDS = ["name", "subject", "status", "creation", "modified", "exp_end_date", "custom_work_category",
    "custom_equipment", "custom_asset", "custom_customer", "custom_service_contract", "custom_started_at",
    "custom_completed_at", "custom_labour_hours", "custom_labour_billing", "custom_labour_mode", "custom_labour_rate",
    "custom_flat_fee", "custom_sales_invoice", "custom_problem_code", "custom_cause_code", "custom_machine_down",
    "custom_down_from", "custom_down_to", "custom_sent_back"]
OPEN = ("Completed", "Cancelled")
REPEAT_DAYS = 30
MIN_SCORED = 3
# Out of 100: what each measure counts towards a technician's score.
WEIGHTS = {"on_time": 30, "first_fix": 35, "no_rework": 20, "coded": 15}


def hours_between(start, end):
    if not start or not end:
        return None
    return max(0.0, (get_datetime(end) - get_datetime(start)).total_seconds() / 3600)


def completed_at(job):
    return job.get("custom_completed_at") or job.get("modified")


def average(values):
    values = [v for v in values if v is not None]
    return round(sum(values) / len(values), 1) if values else None


def buckets(start, end, days):
    """Chart columns covering the period: days for a week, weeks up to a quarter, months for a year."""
    start, end = getdate(start), getdate(end)
    result = []
    while start <= end:
        if days <= 7:
            stop, label = start, start.strftime("%a")
        elif days <= 90:
            stop, label = min(getdate(add_days(start, 6)), end), start.strftime("%d %b")
        else:
            # Calendar months; the first and last may be part months.
            stop = min(getdate(add_days(add_months(start.replace(day=1), 1), -1)), end)
            label = start.strftime("%b %y") if not result else start.strftime("%b")
        result.append({"start": start, "end": stop, "label": label, "created": 0, "completed": 0})
        start = getdate(add_days(stop, 1))
    return result


def pm_state(job, today):
    """On time, Late, Missed or Upcoming for a preventive job against its due date."""
    due = getdate(job.exp_end_date)
    if job.status == "Completed":
        return "On time" if getdate(completed_at(job)) <= due else "Late"
    if job.status == "Cancelled":
        return None
    return "Missed" if due < today else "Upcoming"


def _in(value, start, end):
    return value is not None and getdate(start) <= getdate(value) <= getdate(end)


def _jobs(start):
    """Work orders still open, or created or completed since the start (top-level jobs only)."""
    base = {"custom_work_order": 1, "is_template": 0, "parent_task": ["is", "not set"]}
    names = set(frappe.get_all("Task", filters={**base, "status": ["not in", OPEN]}, pluck="name", limit_page_length=0))
    names |= set(frappe.get_all("Task", filters={**base, "creation": [">=", start]}, pluck="name", limit_page_length=0))
    names |= set(frappe.get_all("Task", filters={**base, "modified": [">=", start]}, pluck="name", limit_page_length=0))
    names |= set(frappe.get_all("Task", filters={**base, "exp_end_date": [">=", start]}, pluck="name", limit_page_length=0))
    # Machines that were down during the period, however old the job.
    names |= set(frappe.get_all("Task", filters={**base, "custom_machine_down": 1}, or_filters=[
        ["custom_down_to", ">=", start], ["custom_down_to", "is", "not set"]], pluck="name", limit_page_length=0))
    if not names:
        return []
    return frappe.get_all("Task", filters={"name": ["in", list(names)]}, fields=JOB_FIELDS, limit_page_length=0)


def _assignees(names):
    result = defaultdict(list)
    for chunk in range(0, len(names), 500):
        for row in frappe.get_all("ToDo", filters={"reference_type": "Task", "reference_name": ["in", names[chunk:chunk + 500]],
                "status": ["!=", "Cancelled"]}, fields=["reference_name", "allocated_to"]):
            if row.allocated_to not in result[row.reference_name]:
                result[row.reference_name].append(row.allocated_to)
    return result


def work_summary(jobs, start, end, today):
    created = [j for j in jobs if _in(j.creation, start, end)]
    done = [j for j in jobs if j.status == "Completed" and _in(completed_at(j), start, end)]
    repairs = [j for j in done if j.custom_work_category == "Corrective"]
    open_now = [j for j in jobs if j.status not in OPEN]
    return {
        "created": len(created), "completed": len(done), "open": len(open_now),
        "overdue": sum(1 for j in open_now if j.exp_end_date and getdate(j.exp_end_date) < today),
        "repair_hours": average(hours_between(j.creation, completed_at(j)) for j in repairs),
        "response_hours": average(hours_between(j.creation, j.custom_started_at) for j in created if j.custom_started_at),
    }


def pm_summary(jobs, start, end, today):
    counts = defaultdict(int)
    late = []
    for job in jobs:
        if job.custom_work_category != "Preventive" or not job.exp_end_date or not _in(job.exp_end_date, start, end):
            continue
        state = pm_state(job, today)
        if state:
            counts[state] += 1
            if state in ("Late", "Missed"):
                late.append({"name": job.name, "subject": job.subject, "due": job.exp_end_date, "state": state})
    judged = counts["On time"] + counts["Late"] + counts["Missed"]
    return {"on_time": counts["On time"], "late": counts["Late"], "missed": counts["Missed"],
        "upcoming": counts["Upcoming"], "compliance": round(100 * counts["On time"] / judged) if judged else None,
        "late_jobs": sorted(late, key=lambda r: str(r["due"]))[:10]}


def by_category(jobs, start, end):
    rows = defaultdict(lambda: {"created": 0, "completed": 0, "hours": []})
    for job in jobs:
        row = rows[job.custom_work_category or "Other"]
        if _in(job.creation, start, end):
            row["created"] += 1
        if job.status == "Completed" and _in(completed_at(job), start, end):
            row["completed"] += 1
            row["hours"].append(hours_between(job.creation, completed_at(job)))
    return sorted(({"category": k, "created": v["created"], "completed": v["completed"],
        "avg_hours": average(v["hours"])} for k, v in rows.items() if v["created"] or v["completed"]),
        key=lambda r: -(r["created"] + r["completed"]))


def machine_of(job):
    return ("equipment", job.custom_equipment) if job.get("custom_equipment") else \
        ("asset", job.custom_asset) if job.get("custom_asset") else None


def repeat_repairs(jobs, days=REPEAT_DAYS):
    """{completed repair: the next repair on the same machine opened within `days` of its completion}."""
    repairs = defaultdict(list)
    for job in jobs:
        key = machine_of(job)
        if key and job.custom_work_category == "Corrective" and job.status != "Cancelled":
            repairs[key].append(job)
    result = {}
    for rows in repairs.values():
        for job in rows:
            if job.status != "Completed":
                continue
            done = get_datetime(completed_at(job))
            later = [o for o in rows if o.name != job.name and done <= get_datetime(o.creation) <= done + timedelta(days=days)]
            if later:
                result[job.name] = min(later, key=lambda o: get_datetime(o.creation))
    return result


def rate(part, whole):
    return round(100 * part / whole) if whole else None


def score(row):
    """Out of 100 from the measures that apply; None until enough jobs are completed."""
    if row["completed"] < MIN_SCORED:
        return None
    measures = {"on_time": row["on_time_rate"], "first_fix": row["first_fix_rate"],
        "no_rework": 100 - rate(row["reopened"], row["completed"]), "coded": row["coded_rate"]}
    used = {k: v for k, v in measures.items() if v is not None}
    return round(sum(WEIGHTS[k] * v for k, v in used.items()) / sum(WEIGHTS[k] for k in used))


def by_technician(jobs, assigned, start, end, today, logged=None, repeats=None):
    """Each person's work in the period and how well it was done (see the scorecard at the top)."""
    logged = logged or {}
    repeats = repeats if repeats is not None else repeat_repairs(jobs)
    rows = defaultdict(lambda: {"completed": 0, "open": 0, "overdue": 0, "labour_hours": 0.0, "hours": [], "due": 0,
        "on_time": 0, "fixes": 0, "first_fix": 0, "repairs": 0, "coded": 0, "reopened": 0, "issues": []})
    for job in jobs:
        done = job.status == "Completed" and _in(completed_at(job), start, end)
        if done and logged.get(job.name):
            for user, hours in logged[job.name].items():
                rows[user]["labour_hours"] += flt(hours)
        for user in assigned.get(job.name, []):
            row = rows[user]
            if done:
                row["completed"] += 1
                if not logged.get(job.name):
                    row["labour_hours"] += flt(job.custom_labour_hours) / len(assigned[job.name])
                row["hours"].append(hours_between(job.creation, completed_at(job)))
                _judge(row, job, repeats)
            elif job.status not in OPEN:
                row["open"] += 1
                if job.exp_end_date and getdate(job.exp_end_date) < today:
                    row["overdue"] += 1
    result = []
    for person in ops.people(list(rows)):
        row = rows[person["user"]]
        entry = {**person, "completed": row["completed"], "open": row["open"], "overdue": row["overdue"],
            "labour_hours": round(row["labour_hours"], 1), "avg_hours": average(row["hours"]),
            "hours_per_job": round(row["labour_hours"] / row["completed"], 1) if row["completed"] and row["labour_hours"] else None,
            "due": row["due"], "on_time": row["on_time"], "on_time_rate": rate(row["on_time"], row["due"]),
            "fixes": row["fixes"], "first_fix": row["first_fix"], "first_fix_rate": rate(row["first_fix"], row["fixes"]),
            "repairs": row["repairs"], "coded": row["coded"], "coded_rate": rate(row["coded"], row["repairs"]),
            "reopened": row["reopened"],
            "issues": sorted(row["issues"], key=lambda i: str(i["on"]), reverse=True)[:15]}
        entry["score"] = score(entry)
        result.append(entry)
    return sorted(result, key=lambda r: (r["score"] is None, -(r["score"] or 0), -r["completed"], -r["open"],
        r["full_name"] or ""))


def _judge(row, job, repeats):
    """Count one completed job towards a person's accuracy measures, noting what went wrong."""
    def issue(kind, detail, other=None):
        row["issues"].append({"kind": kind, "name": job.name, "subject": job.subject, "on": completed_at(job),
            "detail": detail, "other": other})
    if job.exp_end_date:
        row["due"] += 1
        if getdate(completed_at(job)) <= getdate(job.exp_end_date):
            row["on_time"] += 1
        else:
            late = (getdate(completed_at(job)) - getdate(job.exp_end_date)).days
            issue("late", _("Finished {0} day(s) after it was due").format(late))
    if cint(job.get("custom_sent_back")):
        row["reopened"] += 1
        issue("reopened", _("Went back to work {0} time(s) after review or completion").format(cint(job.custom_sent_back)))
    if job.custom_work_category != "Corrective":
        return
    row["repairs"] += 1
    if job.get("custom_problem_code") and job.get("custom_cause_code"):
        row["coded"] += 1
    else:
        issue("uncoded", _("Problem or cause code missing"))
    if machine_of(job):
        row["fixes"] += 1
        again = repeats.get(job.name)
        if again:
            days = max(0, (get_datetime(again.creation) - get_datetime(completed_at(job))).days)
            issue("repeat", _("The machine needed another repair {0} day(s) later").format(days), again.name)
        else:
            row["first_fix"] += 1


def repeat_failures(jobs, start, end):
    """Machines (customer equipment or our own assets) with the most repairs opened in the period."""
    counts = defaultdict(lambda: {"jobs": 0, "hours": []})
    for job in jobs:
        if job.custom_work_category != "Corrective" or not _in(job.creation, start, end):
            continue
        key = ("equipment", job.custom_equipment) if job.custom_equipment else ("asset", job.custom_asset) if job.custom_asset else None
        if key:
            counts[key]["jobs"] += 1
            if job.status == "Completed":
                counts[key]["hours"].append(hours_between(job.creation, completed_at(job)))
    top = sorted(counts.items(), key=lambda kv: -kv[1]["jobs"])[:8]
    from projectit.equipment import labels
    machines = labels(k[1] for k, _v in top if k[0] == "equipment")
    assets = ops.asset_names(k[1] for k, _v in top if k[0] == "asset")
    result = []
    for (kind, name), row in top:
        label = machines.get(name) or {}
        result.append({"kind": kind, "name": name, "repairs": row["jobs"], "avg_hours": average(row["hours"]),
            "label": label.get("equipment_name") if kind == "equipment" else assets.get(name) or name,
            "where": _where(label) if kind == "equipment" else _("Our equipment")})
    return result


def failure_counts(jobs, start, end):
    """The problem and cause codes on jobs completed in the period, most common first."""
    done = [j for j in jobs if j.status == "Completed" and _in(completed_at(j), start, end)]
    used = {j.get(f) for j in done for f in ("custom_problem_code", "custom_cause_code") if j.get(f)}
    titles = dict(frappe.get_all("Techy Failure Code", filters={"name": ["in", list(used)]}, fields=["name", "title"],
        as_list=True)) if used else {}
    result = {}
    for key, field in (("problems", "custom_problem_code"), ("causes", "custom_cause_code")):
        counts = Counter(j.get(field) for j in done if j.get(field))
        result[key] = [{"code": code, "title": titles.get(code, code), "jobs": n} for code, n in counts.most_common(8)]
    repairs = [j for j in done if j.custom_work_category == "Corrective"]
    result.update(repairs=len(repairs), coded=sum(1 for j in repairs if j.get("custom_problem_code")))
    return result


def _window(start, end):
    period_start = get_datetime(getdate(start))
    period_end = min(get_datetime(f"{getdate(end)} 23:59:59"), now_datetime())
    return period_start, period_end, max((period_end - period_start).total_seconds() / 3600, 1)


def _down_by_machine(jobs, start, end):
    """{machine: hours down within the period, stoppages, repair times, down now}."""
    period_start, period_end, _hours = _window(start, end)
    rows = defaultdict(lambda: {"hours": 0.0, "stops": 0, "repairs": [], "down_now": False})
    for job in jobs:
        if not job.get("custom_machine_down") or not job.get("custom_down_from"):
            continue
        key = ("equipment", job.custom_equipment) if job.custom_equipment else ("asset", job.custom_asset) if job.custom_asset else None
        down_from = get_datetime(job.custom_down_from)
        back = get_datetime(job.custom_down_to) if job.get("custom_down_to") else now_datetime()
        overlap = (min(back, period_end) - max(down_from, period_start)).total_seconds() / 3600
        if not key or overlap <= 0:
            continue
        row = rows[key]
        row["hours"] += overlap
        row["stops"] += 1
        if job.get("custom_down_to"):
            row["repairs"].append((back - down_from).total_seconds() / 3600)
        else:
            row["down_now"] = True
    return rows


def downtime(jobs, start, end):
    """Hours each machine was down within the period, its stoppages, time to repair and availability."""
    period_hours = _window(start, end)[2]
    rows = _down_by_machine(jobs, start, end)
    top = sorted(rows.items(), key=lambda kv: -kv[1]["hours"])
    from projectit.equipment import labels
    machines = labels(k[1] for k, _v in top if k[0] == "equipment")
    assets = ops.asset_names(k[1] for k, _v in top if k[0] == "asset")
    result = []
    for (kind, name), row in top[:10]:
        label = machines.get(name) or {}
        result.append({"kind": kind, "name": name, "hours": round(row["hours"], 1), "stops": row["stops"],
            "mttr": average(row["repairs"]), "down_now": row["down_now"],
            "availability": round(max(0.0, 100 * (1 - row["hours"] / period_hours)), 1),
            "label": label.get("equipment_name") if kind == "equipment" else assets.get(name) or name,
            "where": _where(label) if kind == "equipment" else _("Our equipment")})
    repairs = [h for row in rows.values() for h in row["repairs"]]
    return {"machines": result, "hours": round(sum(r["hours"] for r in rows.values()), 1),
        "stops": sum(r["stops"] for r in rows.values()), "mttr": average(repairs),
        "down_now": sum(1 for r in rows.values() if r["down_now"])}


def _machine_info(keys):
    """Name, customer and site for machine keys: customer equipment, or our own assets."""
    equipment = [k[1] for k in keys if k[0] == "equipment"]
    assets = [k[1] for k in keys if k[0] == "asset"]
    info = {}
    for r in frappe.get_all("Techy Equipment", filters={"name": ["in", equipment]}, fields=["name", "equipment_name",
            "customer", "customer_name", "site", "installed_on", "delivered_on"], limit_page_length=0) if equipment else []:
        info[("equipment", r.name)] = {"label": r.equipment_name, "customer": r.customer, "customer_name": r.customer_name,
            "where": _where(r), "since": r.installed_on or r.delivered_on}
    for r in frappe.get_all("Asset", filters={"name": ["in", assets]}, fields=["name", "asset_name",
            "available_for_use_date"], limit_page_length=0) if assets else []:
        info[("asset", r.name)] = {"label": r.asset_name or r.name, "customer": None, "customer_name": _("Our equipment"),
            "where": _("Our equipment"), "since": r.available_for_use_date}
    return info


def parts_cost(names):
    """{job: parts issued to it, at stock valuation cost}."""
    names = list(names)
    if not names:
        return {}
    return {r[0]: flt(r[1]) for r in frappe.db.sql("""SELECT work_order, SUM(qty * valuation_rate)
        FROM `tabTechy Part Usage` WHERE status = 'Issued' AND work_order IN %s GROUP BY work_order""", (tuple(names),))}


def machine_costs(jobs, start, end, logged=None, cost_rate=0):
    """What each machine cost to keep running in the period (parts at cost plus labour), how often it broke down
    and the average running time between breakdowns (MTBF), with the same totals per customer."""
    logged = logged or {}
    period_start, period_end, period_hours = _window(start, end)
    done = [j for j in jobs if j.status == "Completed" and _in(completed_at(j), start, end) and machine_of(j)]
    parts = parts_cost(j.name for j in done)
    down = _down_by_machine(jobs, start, end)
    rows = defaultdict(lambda: {"jobs": 0, "repairs": 0, "parts": 0.0, "hours": 0.0})
    for job in done:
        row = rows[machine_of(job)]
        row["jobs"] += 1
        row["parts"] += parts.get(job.name, 0)
        row["hours"] += sum(logged[job.name].values()) if logged.get(job.name) else flt(job.custom_labour_hours)
    for job in jobs:
        key = machine_of(job)
        if key and job.custom_work_category == "Corrective" and job.status != "Cancelled" and _in(job.creation, start, end):
            rows[key]["repairs"] += 1
    for key in down:
        rows[key]  # machines that were only down still show
    info = _machine_info(list(rows))
    machines = []
    for key, row in rows.items():
        meta = info.get(key) or {"label": key[1], "customer": None, "customer_name": "", "where": "", "since": None}
        # Running time: the period, from when the machine went into service, less the hours it was down.
        since = max(period_start, get_datetime(getdate(meta["since"]))) if meta.get("since") else period_start
        running = max(0.0, (period_end - since).total_seconds() / 3600 - down.get(key, {}).get("hours", 0))
        labour = row["hours"] * flt(cost_rate)
        machines.append({"kind": key[0], "name": key[1], "label": meta["label"], "customer": meta["customer"],
            "customer_name": meta["customer_name"], "where": meta["where"], "jobs": row["jobs"], "breakdowns": row["repairs"],
            "parts_cost": round(row["parts"], 2), "labour_hours": round(row["hours"], 1), "labour_cost": round(labour, 2),
            "cost": round(row["parts"] + labour, 2), "down_hours": round(down.get(key, {}).get("hours", 0), 1),
            "mtbf_hours": round(running / row["repairs"], 1) if row["repairs"] else None})
    machines.sort(key=lambda m: (-m["cost"], -m["breakdowns"], m["label"] or ""))
    customers = defaultdict(lambda: {"machines": 0, "jobs": 0, "breakdowns": 0, "parts_cost": 0.0, "labour_cost": 0.0,
        "cost": 0.0, "down_hours": 0.0})
    for m in machines:
        c = customers[(m["customer"], m["customer_name"])]
        c["machines"] += 1
        for k in ("jobs", "breakdowns", "parts_cost", "labour_cost", "cost", "down_hours"):
            c[k] += m[k]
    by_customer = sorted(({"customer": k[0], "customer_name": k[1], **{f: round(v, 2) if isinstance(v, float) else v
        for f, v in c.items()}} for k, c in customers.items()), key=lambda c: -c["cost"])
    return {"machines": machines, "customers": by_customer, "labour_cost_rate": flt(cost_rate),
        "total": round(sum(m["cost"] for m in machines), 2), "breakdowns": sum(m["breakdowns"] for m in machines),
        "mtbf_hours": average(m["mtbf_hours"] for m in machines)}


def pm_by_customer(jobs, start, end, today):
    """Preventive maintenance due in the period, done on time or not, per customer and per machine."""
    rows = defaultdict(lambda: defaultdict(lambda: {"on_time": 0, "late": 0, "missed": 0, "upcoming": 0}))
    for job in jobs:
        if job.custom_work_category != "Preventive" or not job.exp_end_date or not _in(job.exp_end_date, start, end):
            continue
        state = pm_state(job, today)
        if not state:
            continue
        key = machine_of(job) or ("none", "")
        customer = job.get("custom_customer") or ("" if key[0] != "asset" else None)
        rows[customer][key][state.lower().replace(" ", "_")] += 1
    info = _machine_info([k for c in rows.values() for k in c if k[0] != "none"])
    names = dict(frappe.get_all("Customer", filters={"name": ["in", [c for c in rows if c]]},
        fields=["name", "customer_name"], as_list=True)) if any(rows) else {}

    def compliance(r):
        judged = r["on_time"] + r["late"] + r["missed"]
        return round(100 * r["on_time"] / judged) if judged else None

    result = []
    for customer, machines in rows.items():
        total = {"on_time": 0, "late": 0, "missed": 0, "upcoming": 0}
        listed = []
        for key, r in machines.items():
            for k in total:
                total[k] += r[k]
            meta = info.get(key) or {}
            listed.append({"kind": key[0], "name": key[1], "label": meta.get("label") or _("No machine"),
                "where": meta.get("where") or "", **r, "compliance": compliance(r)})
        listed.sort(key=lambda m: (m["compliance"] if m["compliance"] is not None else 101, m["label"]))
        label = names.get(customer, customer) if customer else _("Our equipment") if customer is None else _("No customer")
        result.append({"customer": customer, "customer_name": label, **total, "compliance": compliance(total),
            "machines": listed})
    return sorted(result, key=lambda c: (c["compliance"] if c["compliance"] is not None else 101, c["customer_name"]))


def cover_ending(days=90):
    """Warranties and service contracts ending in the next `days` days: renewals to offer."""
    today = getdate(nowdate())
    until = getdate(add_days(today, days))
    from projectit.equipment import RETIRED
    rows = []
    for kind, field in (("Warranty", "warranty_until"), ("Contract", "contract_until")):
        for r in frappe.get_all("Techy Equipment", filters={field: ["between", [today, until]],
                "status": ["not in", list(RETIRED)]}, fields=["name", "equipment_name", "customer", "customer_name",
                "site", "service_contract", field], limit_page_length=0):
            ends = getdate(r[field])
            rows.append({"kind": kind, "name": r.name, "label": r.equipment_name, "customer": r.customer,
                "customer_name": r.customer_name, "where": _where(r), "contract": r.service_contract if kind == "Contract" else None,
                "ends": ends, "days_left": (ends - today).days})
    contracts_ending = []
    for c in frappe.get_all("Techy Service Contract", filters={"status": ["!=", "Cancelled"],
            "end_date": ["between", [today, until]]}, fields=["name", "title", "customer", "customer_name", "end_date"],
            limit_page_length=0):
        machines = frappe.db.count("Techy Contract Equipment", {"parent": c.name, "parenttype": "Techy Service Contract"})
        contracts_ending.append({"name": c.name, "title": c.title, "customer": c.customer, "customer_name": c.customer_name,
            "ends": getdate(c.end_date), "days_left": (getdate(c.end_date) - today).days, "machines": machines})
    return {"days": days, "machines": sorted(rows, key=lambda r: (r["ends"], r["customer_name"] or "")),
        "contracts": sorted(contracts_ending, key=lambda r: r["ends"])}


def _where(label):
    customer, site = label.get("customer_name"), label.get("site")
    if site and customer and customer in site:
        return site
    return ", ".join(x for x in (customer, site) if x)


def request_summary(start, end):
    rows = frappe.get_all("Issue", filters={"custom_maintenance_request": 1, "creation": [">=", start]},
        fields=["status", "custom_work_order", "creation"], limit_page_length=0)
    rows = [r for r in rows if _in(r.creation, start, end)]
    closed = ("Resolved", "Closed")
    return {"received": len(rows), "approved": sum(1 for r in rows if r.custom_work_order),
        "declined": sum(1 for r in rows if not r.custom_work_order and r.status in closed),
        "pending": sum(1 for r in rows if not r.custom_work_order and r.status not in closed)}


def contract_profit(start, end, jobs, logged=None):
    """Money in and out for each service contract that ran during the period."""
    logged = logged or {}
    cost_rate = flt(frappe.get_cached_doc("Techy Ops Settings").get("labour_cost_rate"))
    rows = frappe.get_all("Techy Service Contract", fields=["name", "title", "customer_name", "status", "subscription",
        "start_date", "end_date"], limit_page_length=0)
    result = []
    for doc in rows:
        if doc.status == "Cancelled" or (doc.start_date and getdate(doc.start_date) > getdate(end)) \
                or (doc.end_date and getdate(doc.end_date) < getdate(start)):
            continue
        machines = set(frappe.get_all("Techy Contract Equipment", filters={"parent": doc.name,
            "parenttype": "Techy Service Contract"}, pluck="equipment"))
        own = [j for j in jobs if j.status == "Completed" and _in(completed_at(j), start, end)
            and (j.custom_service_contract == doc.name or (j.custom_equipment and j.custom_equipment in machines))]
        subscription = 0
        if doc.subscription:
            subscription = flt(frappe.db.sql("""SELECT SUM(base_net_total) FROM `tabSales Invoice`
                WHERE subscription = %s AND docstatus = 1 AND posting_date BETWEEN %s AND %s""",
                (doc.subscription, getdate(start), getdate(end)))[0][0])
        invoiced = {name for name in frappe.get_all("Sales Invoice", filters={"docstatus": 1,
            "name": ["in", [j.custom_sales_invoice for j in own if j.custom_sales_invoice] or [""]]}, pluck="name")}
        extra = 0
        for job in own:
            if job.custom_sales_invoice in invoiced:
                labour, part_total = billing._job_total(job)
                extra += labour["amount"] + part_total
        parts_cost = flt(frappe.db.sql("""SELECT SUM(qty * valuation_rate) FROM `tabTechy Part Usage`
            WHERE status = 'Issued' AND work_order IN %s""", (tuple(j.name for j in own) or ("",),))[0][0])
        hours = sum(sum(logged[j.name].values()) if logged.get(j.name) else flt(j.custom_labour_hours) for j in own)
        income, cost = subscription + extra, parts_cost + hours * cost_rate
        result.append({"name": doc.name, "title": doc.title, "customer_name": doc.customer_name,
            "state": contracts.state(doc), "jobs": len(own), "subscription": subscription, "extra": extra,
            "income": income, "parts_cost": parts_cost, "labour_hours": round(hours, 1), "labour_cost": hours * cost_rate,
            "cost": cost, "profit": income - cost, "margin": round(100 * (income - cost) / income) if income else None})
    return {"rows": sorted(result, key=lambda r: r["profit"]), "labour_cost_rate": cost_rate}


def _period(days):
    days = cint(days) if cint(days) in PERIODS else 30
    today = getdate(nowdate())
    start = getdate(add_days(today, 1 - days))
    return days, today, start, today, getdate(add_days(start, -days))


@frappe.whitelist()
def machines(days=30, cover_days=90):
    """The Machines tab: cost and reliability per machine and customer, PM on time per customer, cover ending."""
    ops.system_user()
    ops.require_manager()
    days, today, start, end, before = _period(days)
    jobs = _jobs(before)
    logged = timer.hours_by_person([j.name for j in jobs if j.status == "Completed"])
    rate = flt(frappe.get_cached_doc("Techy Ops Settings").get("labour_cost_rate"))
    return {"days": days, "start": start, "end": end, "currency": frappe.db.get_default("currency"),
        "costs": machine_costs(jobs, start, end, logged, rate), "pm": pm_by_customer(jobs, start, end, today),
        "cover": cover_ending(cint(cover_days) if cint(cover_days) in (30, 60, 90) else 90)}


@frappe.whitelist()
def export(days=30):
    """Every report table for the period as one Excel workbook, a sheet per table."""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter
    r = dashboard(days)
    m = machines(days)
    period = f"{r['start']} to {r['end']}"
    pct = lambda v: f"{v}%" if v is not None else ""
    sheets = [
        (_("Summary"), [_("Measure"), _("Value")], [
            [_("Period"), period], [_("Jobs completed"), r["work"]["completed"]], [_("Jobs opened"), r["work"]["created"]],
            [_("Open now"), r["work"]["open"]], [_("Overdue"), r["work"]["overdue"]],
            [_("Avg hours to fix repairs"), r["work"]["repair_hours"]], [_("Avg hours to start"), r["work"]["response_hours"]],
            [_("PM done on time"), pct(r["pm"]["compliance"])], [_("Requests received"), r["requests"]["received"]],
            [_("Machine downtime (hours)"), r["downtime"]["hours"]], [_("Stoppages"), r["downtime"]["stops"]],
            [_("Machine costs"), m["costs"]["total"]], [_("Breakdowns"), m["costs"]["breakdowns"]]]),
        (_("Technicians"), [_("Person"), _("Score"), _("Completed"), _("On time %"), _("Fixed first time %"), _("Sent back"),
            _("Coded %"), _("Avg hours to finish"), _("Hours logged"), _("Hours per job"), _("Open"), _("Overdue")],
            [[t["full_name"], t["score"], t["completed"], t["on_time_rate"], t["first_fix_rate"], t["reopened"], t["coded_rate"],
                t["avg_hours"], t["labour_hours"], t["hours_per_job"], t["open"], t["overdue"]] for t in r["technicians"]]),
        (_("Machine costs"), [_("Machine"), _("Customer / site"), _("Jobs"), _("Breakdowns"), _("MTBF hours"), _("Hours down"),
            _("Parts cost"), _("Labour hours"), _("Labour cost"), _("Total cost")],
            [[x["label"], x["where"], x["jobs"], x["breakdowns"], x["mtbf_hours"], x["down_hours"], x["parts_cost"],
                x["labour_hours"], x["labour_cost"], x["cost"]] for x in m["costs"]["machines"]]),
        (_("Customer costs"), [_("Customer"), _("Machines"), _("Jobs"), _("Breakdowns"), _("Hours down"), _("Parts cost"),
            _("Labour cost"), _("Total cost")],
            [[c["customer_name"], c["machines"], c["jobs"], c["breakdowns"], c["down_hours"], c["parts_cost"], c["labour_cost"],
                c["cost"]] for c in m["costs"]["customers"]]),
        (_("PM by customer"), [_("Customer"), _("Machine"), _("On time"), _("Late"), _("Missed"), _("Upcoming"), _("On time %")],
            [[c["customer_name"], x["label"], x["on_time"], x["late"], x["missed"], x["upcoming"], x["compliance"]]
                for c in m["pm"] for x in c["machines"]]),
        (_("Downtime"), [_("Machine"), _("Where"), _("Hours down"), _("Stoppages"), _("Avg hours to repair"), _("Availability %"),
            _("Down now")], [[x["label"], x["where"], x["hours"], x["stops"], x["mttr"], x["availability"],
                _("Yes") if x["down_now"] else ""] for x in r["downtime"]["machines"]]),
        (_("Most repairs"), [_("Machine"), _("Where"), _("Repairs"), _("Avg hours to fix")],
            [[x["label"], x["where"], x["repairs"], x["avg_hours"]] for x in r["machines"]]),
        (_("Failures"), [_("Kind"), _("Code"), _("Title"), _("Jobs")],
            [[_("Problem"), x["code"], x["title"], x["jobs"]] for x in r["failures"]["problems"]]
            + [[_("Cause"), x["code"], x["title"], x["jobs"]] for x in r["failures"]["causes"]]),
        (_("Cover ending"), [_("Cover"), _("Machine"), _("Customer"), _("Ends"), _("Days left")],
            [[x["kind"], x["label"], x["customer_name"], x["ends"], x["days_left"]] for x in m["cover"]["machines"]]
            + [[_("Contract"), c["title"], c["customer_name"], c["ends"], c["days_left"]] for c in m["cover"]["contracts"]]),
        (_("Contract profit"), [_("Contract"), _("Customer"), _("Jobs"), _("Subscription"), _("Extra billed"), _("Parts cost"),
            _("Labour cost"), _("Profit"), _("Margin %")],
            [[c["title"], c["customer_name"], c["jobs"], c["subscription"], c["extra"], c["parts_cost"], c["labour_cost"],
                c["profit"], c["margin"]] for c in r["contracts"]["rows"]]),
    ]
    book = Workbook()
    book.remove(book.active)
    for title, header, rows in sheets:
        sheet = book.create_sheet(title[:31])
        sheet.append(header)
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        for row in rows:
            sheet.append([v if v is None or isinstance(v, (int, float, str)) else str(v) for v in row])
        for i, name in enumerate(header, 1):
            width = max([len(str(name))] + [len(str(row[i - 1] or "")) for row in rows]) + 2
            sheet.column_dimensions[get_column_letter(i)].width = min(width, 50)
        sheet.freeze_panes = "A2"
    import io
    out = io.BytesIO()
    book.save(out)
    frappe.local.response.filename = f"Techy reports {r['start']} to {r['end']}.xlsx"
    frappe.local.response.filecontent = out.getvalue()
    frappe.local.response.type = "download"


@frappe.whitelist()
def dashboard(days=30):
    ops.system_user()
    ops.require_manager()
    days = cint(days) if cint(days) in PERIODS else 30
    today = getdate(nowdate())
    start, end = getdate(add_days(today, 1 - days)), today
    before = getdate(add_days(start, -days))
    jobs = _jobs(before)
    assigned = _assignees([j.name for j in jobs])
    logged = timer.hours_by_person([j.name for j in jobs if j.status == "Completed"])
    chart = buckets(start, end, days)
    for job in jobs:
        for col in chart:
            if _in(job.creation, col["start"], col["end"]):
                col["created"] += 1
            if job.status == "Completed" and _in(completed_at(job), col["start"], col["end"]):
                col["completed"] += 1
    return {
        "days": days, "start": start, "end": end, "currency": frappe.db.get_default("currency"),
        "work": work_summary(jobs, start, end, today),
        "previous": work_summary(jobs, before, getdate(add_days(start, -1)), getdate(add_days(start, -1))),
        "chart": [{k: c[k] for k in ("label", "created", "completed")} for c in chart],
        "pm": pm_summary(jobs, start, end, today),
        "categories": by_category(jobs, start, end),
        "technicians": by_technician(jobs, assigned, start, end, today, logged),
        "scoring": {"weights": WEIGHTS, "min_jobs": MIN_SCORED, "repeat_days": REPEAT_DAYS},
        "machines": repeat_failures(jobs, start, end),
        "requests": request_summary(start, end),
        "contracts": contract_profit(start, end, jobs, logged),
        "failures": failure_counts(jobs, start, end),
        "downtime": downtime(jobs, start, end),
    }
