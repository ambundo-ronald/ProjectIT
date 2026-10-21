"""Staff tools for the customer app: every customer account, and how well each machine is being logged.

Managers only. The same numbers as the Health tab's "How complete the data is" card (health.customer_logging).
"""
import io

import frappe
from frappe import _
from frappe.utils import cint, get_fullname

from projectit import ops

DAYS = (7, 30, 90)


def _days(value):
    return cint(value) if cint(value) in DAYS else 30


def _match(text, *values):
    text = (text or "").strip().lower()
    return not text or any(text in str(v or "").lower() for v in values)


def logging_rows(days=30, search=None):
    """One row per customer machine someone can log: schedule, logged of expected, last log, worst first."""
    from projectit import customer_notify, health
    from projectit.plant import LOG
    customers = frappe.get_all("Techy Customer User", filters={"enabled": 1}, pluck="customer", distinct=True)
    if not customers:
        return []
    machines = frappe.get_all("Techy Equipment", filters={"customer": ["in", customers], "status": "Active"},
        fields=["name", "equipment_name", "customer", "customer_name", "site", "log_schedule"], order_by="customer_name asc")
    rows = []
    for m in machines:
        if not _match(search, m.equipment_name, m.customer_name, m.site, m.name):
            continue
        logging = health.customer_logging(m.name, days)
        if logging is None:
            continue  # nobody at the customer can see this machine
        last = frappe.db.get_value(LOG, {"equipment": m.name}, ["log_on", "logged_by"], order_by="log_on desc", as_dict=True)
        share = round(100 * logging["done"] / logging["expected"]) if logging["expected"] else None
        rows.append({"equipment": m.name, "equipment_name": m.equipment_name, "customer": m.customer,
            "customer_name": m.customer_name, "site": m.site, "schedule": logging["schedule"], "expected": logging["expected"],
            "done": logging["done"], "missed": max(0, logging["expected"] - logging["done"]), "logs": logging["logs"],
            "share": share, "last_on": last.log_on if last else None,
            "last_by": get_fullname(last.logged_by) if last else None,
            "people": len(customer_notify.audience(m.name))})
    # Worst first: machines that should be logged and are not; "not expected" last.
    return sorted(rows, key=lambda r: (r["share"] is None, r["share"] if r["share"] is not None else 0, r["customer_name"] or ""))


@frappe.whitelist()
def logging(days=30, search=None):
    ops.system_user()
    ops.require_manager()
    days = _days(days)
    rows = logging_rows(days, search)
    expected = [r for r in rows if r["expected"]]
    return {"days": days, "rows": rows, "machines": len(rows),
        "good": sum(1 for r in expected if r["share"] >= 90), "some": sum(1 for r in expected if 0 < r["share"] < 90),
        "none": sum(1 for r in expected if r["share"] == 0),
        "overall": round(100 * sum(r["done"] for r in expected) / sum(r["expected"] for r in expected)) if expected else None}


@frappe.whitelist()
def accounts(search=None):
    """Every customer account, with when they last signed in and logged."""
    ops.system_user()
    ops.require_manager()
    from projectit.customer_accounts import _row
    from projectit.plant import LOG
    names = frappe.get_all("Techy Customer User", fields=["name", "customer_name", "full_name", "user"],
        order_by="customer_name asc, full_name asc")
    rows = []
    for n in names:
        if not _match(search, n.full_name, n.user, n.customer_name):
            continue
        doc = frappe.get_doc("Techy Customer User", n.name)
        row = _row(doc)
        row.update(customer=doc.customer, customer_name=doc.customer_name or doc.customer, me=False,
            last_log=frappe.db.get_value(LOG, {"logged_by": doc.user}, "log_on", order_by="log_on desc"))
        rows.append(row)
    return {"accounts": rows, "customers": len({r["customer"] for r in rows}),
        "joined": sum(1 for r in rows if r["joined"] and r["enabled"]), "waiting": sum(1 for r in rows if not r["joined"] and r["enabled"]),
        "off": sum(1 for r in rows if not r["enabled"])}


@frappe.whitelist()
def export(days=30):
    """The logging report and the accounts as an Excel workbook."""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter
    report = logging(days)
    people = accounts()
    sheets = [
        (_("Logging"), [_("Customer"), _("Machine"), _("Site"), _("Schedule"), _("Expected"), _("Logged"), _("Missed"),
            _("Logged %"), _("Last log"), _("Last logged by"), _("People")],
            [[r["customer_name"], r["equipment_name"], r["site"], r["schedule"], r["expected"], r["done"], r["missed"],
                r["share"], str(r["last_on"] or ""), r["last_by"], r["people"]] for r in report["rows"]]),
        (_("Accounts"), [_("Customer"), _("Name"), _("Email"), _("Access"), _("Machines"), _("Status"), _("Last sign-in"),
            _("Last log")],
            [[a["customer_name"], a["full_name"], a["user"], a["access"],
                _("All") if a["access"] == "Customer admin" else ", ".join(a["machine_names"]),
                _("Switched off") if not a["enabled"] else _("Active") if a["joined"] else _("Invited"),
                str(a["last_login"] or ""), str(a["last_log"] or "")] for a in people["accounts"]]),
    ]
    book = Workbook()
    book.remove(book.active)
    for title, header, rows in sheets:
        sheet = book.create_sheet(title)
        sheet.append(header)
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        for row in rows:
            sheet.append(row)
        for i, name in enumerate(header, 1):
            width = max([len(str(name))] + [len(str(row[i - 1] or "")) for row in rows]) + 2
            sheet.column_dimensions[get_column_letter(i)].width = min(width, 50)
        sheet.freeze_panes = "A2"
    out = io.BytesIO()
    book.save(out)
    frappe.local.response.filename = f"Customer app logging, last {report['days']} days.xlsx"
    frappe.local.response.filecontent = out.getvalue()
    frappe.local.response.type = "download"


# Every customer plant's status -----------------------------------------------------------------------------------

ORDER = {"Action needed": 0, "Needs attention": 1, "Running well": 2}


@frappe.whitelist()
def plant_status(level=None, search=None):
    """Every customer plant by status, the worst first, with what is already planned: where to send someone today."""
    import json
    ops.system_user()
    ops.require_manager()
    from projectit import customer_notify
    from projectit.plant import LOG, on_the_way
    machines = frappe.get_all("Techy Equipment", filters={"status": "Active", "customer": ["is", "set"],
        "plant_status": ["is", "set"]}, fields=["name", "equipment_name", "customer", "customer_name", "site", "tag",
        "plant_status", "plant_status_reasons", "plant_status_on"])
    counts = {k: 0 for k in ORDER}
    rows = []
    for m in machines:
        counts[m.plant_status] = counts.get(m.plant_status, 0) + 1
        if level and m.plant_status != level:
            continue
        if not _match(search, m.equipment_name, m.customer_name, m.site, m.tag, m.name):
            continue
        jobs = frappe.get_all("Task", filters={"custom_work_order": 1, "custom_equipment": m.name, "parent_task": ["is", "not set"],
            "is_template": 0, "status": ["not in", ["Completed", "Cancelled"]]}, fields=["name", "status", "exp_start_date",
            "custom_on_way_at"], order_by="exp_start_date asc")
        planned = next((j.exp_start_date for j in jobs if j.exp_start_date), None)
        try:
            reasons = json.loads(m.plant_status_reasons or "[]")
        except ValueError:
            reasons = []
        rows.append({"equipment": m.name, "equipment_name": m.equipment_name, "customer": m.customer,
            "customer_name": m.customer_name, "site": m.site, "status": m.plant_status, "reasons": reasons,
            "checked": m.plant_status_on, "open_jobs": len(jobs), "planned": planned,
            "on_way": any(on_the_way(j) for j in jobs),
            "last_log": frappe.db.get_value(LOG, {"equipment": m.name}, "log_on", order_by="log_on desc"),
            "has_app": bool(customer_notify.audience(m.name))})
    rows.sort(key=lambda r: (ORDER.get(r["status"], 3), r["open_jobs"] > 0, r["customer_name"] or ""))
    return {"rows": rows[:300], "total": len(rows), "counts": counts}
